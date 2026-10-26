"""Read the supplied workbook; never modify it. Fail on unparsed course cells."""
import json
import re
import sys
from pathlib import Path
from datetime import date, timedelta
import openpyxl

root = Path(__file__).resolve().parents[1]
source = Path(sys.argv[1]) if len(sys.argv) > 1 else Path('01课表.xlsx')
sheet = openpyxl.load_workbook(source, data_only=True).worksheets[0]
catalog = []
for row in range(3, 17):
    name = str(sheet.cell(row, 3).value).strip()
    names = re.sub(r'（[^）]*）', '', str(sheet.cell(row, 5).value)).split('、')
    catalog.append(dict(id=str(row-2), name=name, teachers=names, assessment=sheet.cell(row,6).value))
lookup = {c['name']:c for c in catalog}
lookup['大学体育'] = lookup['大学体育（三）']
entries = []
day = None
period = None
for row in range(17, 96):
    a, b = sheet.cell(row,1).value, sheet.cell(row,2).value
    if a:
        day = ['星期一','星期二','星期三','星期四','星期五','星期六','星期日'].index(a)
    if b is not None:
        p = str(b).split('-')
        period = [int(p[0]), int(p[-1])]
    for col in range(3,7):
        cell = sheet.cell(row,col)
        if not cell.value:
            continue
        raw = str(cell.value).strip()
        match = re.fullmatch(r'(.+?)\s+([\d、-]+)周\s+([合1-4FC、]+)\s+(.+)',raw)
        assert match, (cell.coordinate,raw)
        name, week_text, group, location = match.groups()
        weeks = []
        for part in week_text.split('、'):
            endpoints = [int(n) for n in part.split('-')]
            weeks.extend(range(endpoints[0],endpoints[-1]+1))
        large = '合' in group
        classes = [1,2,3,4] if large else sorted(set(int(n) for n in re.findall('[1-4]',group)))
        loc = re.fullmatch(r'(.*?)([A-Z]?\d+)',location)
        building, room = loc.groups() if loc else (location,'')
        course = lookup[name]
        entries.append(dict(id=cell.coordinate, courseId=course['id'], name=name, teachers=course['teachers'],
                            day=day,start=period[0],end=period[1],weeks=sorted(set(weeks)),classes=classes,
                            kind='大课' if large else '小课',building=building,room=room,
                            sourceCell=cell.coordinate,raw=raw))

exams=[]
for c in catalog:
    m=re.fullmatch(r'(\d+)周周(.)（(\d+)年(\d+)月(\d+)日）(下午|晚上)',c['assessment'])
    if m:
        week, weekday, year, month, day_of_month, time = m.groups()
        examdate=date(int(year),int(month),int(day_of_month))
        assert examdate-timedelta(weeks=int(week)-1,days='一二三四五六日'.index(weekday)) == date(2026,8,31)
        exams.append(dict(courseId=c['id'],name=c['name'],week=int(week),date=examdate.isoformat(),time=time))

data=dict(title=sheet['A1'].value,semesterStart='2026-08-31',weekCount=19,classes=[1,2,3,4],
          source='01课表.xlsx',teacherNote='教师为原课表列出的课程教师名单，原表未指定每次授课教师。',
          timeSource='节次时间参照用户提供的课表截图。',
          periods=[['08:00','08:45'],['08:55','09:40'],['10:10','10:55'],['11:05','11:50'],
                   ['14:00','14:45'],['14:50','15:35'],['15:55','16:40'],['16:45','17:30'],
                   ['18:30','19:15'],['19:20','20:05'],['20:15','21:00'],['21:05','21:50']],
          courses=catalog,entries=entries,exams=exams)
(root/'web'/'schedule.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
conflicts=[]
for week in range(1,20):
    for cls in range(1,5):
        slots={}
        for e in entries:
            if week in e['weeks'] and cls in e['classes']:
                for period in range(e['start'],e['end']+1):
                    key=(e['day'],period)
                    if key in slots:
                        conflicts.append(dict(week=week,cls=cls,day=e['day'],period=period,cells=[slots[key],e['id']]))
                    slots[key]=e['id']
print(json.dumps(dict(courses=len(catalog),sourceEntries=len(entries),exams=len(exams),conflicts=conflicts),ensure_ascii=False))
